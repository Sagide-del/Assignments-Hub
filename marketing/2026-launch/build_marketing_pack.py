from __future__ import annotations

import io
import os
from pathlib import Path
from textwrap import wrap

import qrcode
from PIL import Image, ImageChops, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.graphics import renderPDF
from reportlab.graphics.barcode import qr
from reportlab.graphics.shapes import Drawing
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
TMP = OUT / "tmp"
TMP.mkdir(parents=True, exist_ok=True)

LOGO = ROOT / "frontend" / "public" / "logo.png"
HERO = ROOT / "frontend" / "public" / "classroom-environment.png"

NAVY = "#101820"
LIME = "#B5E61D"
BACKGROUND = "#F8FAFC"
SLATE = "#475569"
MUTED = "#94A3B8"
LINE = "#D9E1EA"
WHITE = "#FFFFFF"

WEBSITE = "assignmenthub.co.ke"
REGISTER_URL = "https://assignmenthub.co.ke/register"
PHONE = "0748 519 923"
TILL = "4382411"
STORE = "4840093"

FONT_REGULAR = Path("C:/Windows/Fonts/segoeui.ttf")
FONT_SEMIBOLD = Path("C:/Windows/Fonts/seguisb.ttf")
FONT_BOLD = Path("C:/Windows/Fonts/segoeuib.ttf")


def pil_font(size: int, bold: bool = False, semibold: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_BOLD if bold else FONT_SEMIBOLD if semibold else FONT_REGULAR
    return ImageFont.truetype(str(path), size=size)


def cropped_logo() -> Image.Image:
    image = Image.open(LOGO).convert("RGBA")
    rgb = image.convert("RGB")
    background = Image.new("RGB", rgb.size, "white")
    difference = ImageChops.difference(rgb, background).convert("L")
    difference = difference.point(lambda value: 255 if value > 12 else 0)
    bbox = difference.getbbox()
    if not bbox:
        return image
    left, top, right, bottom = bbox
    padding = 18
    return image.crop((max(0, left - padding), max(0, top - padding), min(image.width, right + padding), min(image.height, bottom + padding)))


def logo_reader() -> ImageReader:
    stream = io.BytesIO()
    cropped_logo().save(stream, format="PNG")
    stream.seek(0)
    return ImageReader(stream)


def draw_pdf_logo(pdf: canvas.Canvas, x: float, y: float, width: float) -> None:
    logo = cropped_logo()
    height = width * logo.height / logo.width
    pdf.drawImage(logo_reader(), x, y, width=width, height=height, mask="auto")


def draw_qr_pdf(pdf: canvas.Canvas, value: str, x: float, y: float, size: float) -> None:
    widget = qr.QrCodeWidget(value)
    bounds = widget.getBounds()
    width = bounds[2] - bounds[0]
    height = bounds[3] - bounds[1]
    drawing = Drawing(size, size, transform=[size / width, 0, 0, size / height, 0, 0])
    drawing.add(widget)
    renderPDF.draw(drawing, pdf, x, y)


def qr_png(value: str, output: Path, size: int = 500) -> None:
    code = qrcode.QRCode(version=None, box_size=10, border=2)
    code.add_data(value)
    code.make(fit=True)
    image = code.make_image(fill_color=NAVY, back_color=WHITE).convert("RGB")
    image.resize((size, size), Image.Resampling.NEAREST).save(output)


def fit_image(source: Path, size: tuple[int, int], output: Path, position: tuple[float, float] = (0.5, 0.5)) -> None:
    image = Image.open(source).convert("RGB")
    fitted = ImageOps.fit(image, size, method=Image.Resampling.LANCZOS, centering=position)
    fitted.save(output, quality=94)


def wrapped_lines(text: str, font_name: str, font_size: float, max_width: float) -> list[str]:
    words = text.split()
    lines: list[str] = []
    line: list[str] = []
    for word in words:
        candidate = " ".join([*line, word])
        if line and stringWidth(candidate, font_name, font_size) > max_width:
            lines.append(" ".join(line))
            line = [word]
        else:
            line.append(word)
    if line:
        lines.append(" ".join(line))
    return lines


def draw_pdf_text(pdf: canvas.Canvas, text: str, x: float, y: float, max_width: float, font: str, size: float, color: str, leading: float | None = None) -> float:
    leading = leading or size * 1.25
    pdf.setFont(font, size)
    pdf.setFillColor(HexColor(color))
    for line in wrapped_lines(text, font, size, max_width):
        pdf.drawString(x, y, line)
        y -= leading
    return y


def build_student_pamphlet() -> Path:
    output = OUT / "Assignment-Hub-Student-Pamphlet-A4.pdf"
    width, height = A4
    hero = TMP / "pamphlet-hero.jpg"
    fit_image(HERO, (1800, 1200), hero, position=(0.68, 0.5))

    pdf = canvas.Canvas(str(output), pagesize=A4)
    pdf.setTitle("Assignment Hub Student Pamphlet")
    pdf.setAuthor("Andolih EdTech Studios")

    # Front page
    pdf.setFillColor(HexColor(NAVY))
    pdf.rect(0, 0, width, height, fill=1, stroke=0)
    pdf.drawImage(str(hero), 0, height * 0.44, width=width, height=height * 0.56, mask="auto")
    pdf.setFillColor(HexColor(NAVY))
    pdf.setFillAlpha(0.28)
    pdf.rect(0, height * 0.44, width, height * 0.56, fill=1, stroke=0)
    pdf.setFillAlpha(1)

    pdf.setFillColor(white)
    pdf.roundRect(34, height - 82, 238, 42, 12, fill=1, stroke=0)
    draw_pdf_logo(pdf, 46, height - 74, 210)

    pdf.setFillColor(HexColor(LIME))
    pdf.roundRect(34, 323, 64, 5, 2, fill=1, stroke=0)
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(34, 294, "LEARN. PRACTISE. IMPROVE.")
    y = draw_pdf_text(pdf, "Your learning, all in one place.", 34, 257, width - 68, "Helvetica-Bold", 31, WHITE, 37)
    y = draw_pdf_text(pdf, "Complete assignments, get feedback, practise weak topics and build practical skills from your phone or computer.", 34, y - 8, width - 86, "Helvetica", 13.5, "#DCE5EC", 19)

    pdf.setFillColor(HexColor(LIME))
    pdf.roundRect(34, 92, width - 68, 76, 18, fill=1, stroke=0)
    pdf.setFillColor(HexColor(NAVY))
    pdf.setFont("Helvetica-Bold", 17)
    pdf.drawString(54, 135, "Independent learner access: KSh 300 per month")
    pdf.setFont("Helvetica", 11.5)
    pdf.drawString(54, 112, "Register at assignmenthub.co.ke/register")

    pdf.setFillColor(white)
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(34, 54, WEBSITE)
    pdf.setFont("Helvetica", 12)
    pdf.drawRightString(width - 34, 54, f"WhatsApp {PHONE}")
    pdf.showPage()

    # Back page
    pdf.setFillColor(HexColor(BACKGROUND))
    pdf.rect(0, 0, width, height, fill=1, stroke=0)
    draw_pdf_logo(pdf, 38, height - 70, 190)
    pdf.setFillColor(HexColor(NAVY))
    pdf.setFont("Helvetica-Bold", 25)
    pdf.drawString(38, height - 116, "Everything you need to keep learning")

    features = [
        ("01", "Assignments", "Complete school or independent assignments and submit from any device."),
        ("02", "Instant feedback", "See scores and explanations after results are released."),
        ("03", "STEM Labs", "Follow practical science activities with media, steps and questions."),
        ("04", "Practical skills", "Learn Python, robotics, first aid, machine learning and research skills."),
        ("05", "Study resources", "Access mnemonic cards, revision materials and downloadable learning resources."),
        ("06", "Personal study plan", "Focus study time on subjects and topics that need the most attention."),
    ]
    start_y = height - 160
    card_w = (width - 94) / 2
    card_h = 105
    for index, (number, title, description) in enumerate(features):
        column = index % 2
        row = index // 2
        x = 38 + column * (card_w + 18)
        y_card = start_y - row * (card_h + 15) - card_h
        pdf.setFillColor(white)
        pdf.setStrokeColor(HexColor(LINE))
        pdf.roundRect(x, y_card, card_w, card_h, 16, fill=1, stroke=1)
        pdf.setFillColor(HexColor(LIME))
        pdf.roundRect(x + 14, y_card + card_h - 40, 30, 24, 8, fill=1, stroke=0)
        pdf.setFillColor(HexColor(NAVY))
        pdf.setFont("Helvetica-Bold", 9)
        pdf.drawCentredString(x + 29, y_card + card_h - 32, number)
        pdf.setFont("Helvetica-Bold", 13)
        pdf.drawString(x + 52, y_card + card_h - 33, title)
        draw_pdf_text(pdf, description, x + 15, y_card + card_h - 60, card_w - 30, "Helvetica", 9.4, SLATE, 13)

    pdf.setFillColor(HexColor(NAVY))
    pdf.roundRect(38, 80, width - 76, 145, 20, fill=1, stroke=0)
    pdf.setFillColor(HexColor(LIME))
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(58, 195, "GET STARTED")
    pdf.setFillColor(white)
    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawString(58, 165, "Register, pay and start learning")
    pdf.setFont("Helvetica", 10.5)
    pdf.drawString(58, 137, f"M-Pesa Till {TILL}  |  Store {STORE}")
    pdf.drawString(58, 117, f"Support: {PHONE}  |  {WEBSITE}")
    pdf.setFillColor(white)
    pdf.roundRect(width - 152, 94, 116, 124, 12, fill=1, stroke=0)
    draw_qr_pdf(pdf, REGISTER_URL, width - 143, 113, 98)
    pdf.setFillColor(HexColor(NAVY))
    pdf.setFont("Helvetica", 7.5)
    pdf.drawCentredString(width - 94, 103, "Scan to register")

    pdf.setFillColor(HexColor(SLATE))
    pdf.setFont("Helvetica", 8.5)
    pdf.drawString(38, 42, "Assignment Hub is a web-based learning platform. Internet access is required.")
    pdf.drawRightString(width - 38, 42, "Powered by Andolih EdTech Studios")
    pdf.save()
    return output


def build_nyangorri_snapshot() -> Path:
    output = OUT / "Assignment-Hub-Nyangorri-Boys-Product-Snapshot.pdf"
    width, height = A4
    pdf = canvas.Canvas(str(output), pagesize=A4)
    pdf.setTitle("Assignment Hub Product Snapshot for Nyangorri Boys School")
    pdf.setAuthor("Andolih EdTech Studios")

    pdf.setFillColor(HexColor(BACKGROUND))
    pdf.rect(0, 0, width, height, fill=1, stroke=0)
    pdf.setFillColor(white)
    pdf.rect(0, height - 94, width, 94, fill=1, stroke=0)
    draw_pdf_logo(pdf, 38, height - 68, 194)
    pdf.setFillColor(HexColor(SLATE))
    pdf.setFont("Helvetica-Bold", 9)
    pdf.drawRightString(width - 38, height - 47, "PRODUCT SNAPSHOT")
    pdf.setFont("Helvetica", 9)
    pdf.drawRightString(width - 38, height - 63, "Prepared for the Nyangorri Boys School community")

    pdf.setFillColor(HexColor(NAVY))
    pdf.roundRect(38, height - 260, width - 76, 132, 22, fill=1, stroke=0)
    pdf.setFillColor(HexColor(LIME))
    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(60, height - 161, "A CONNECTED LEARNING WORKSPACE")
    draw_pdf_text(pdf, "Assignments, feedback, resources and practical skills on any device.", 60, height - 191, width - 150, "Helvetica-Bold", 20.5, WHITE, 27)
    pdf.setFillColor(HexColor("#C9D4DC"))
    pdf.setFont("Helvetica", 10.5)
    pdf.drawString(60, height - 237, "No app installation required. Learners sign in through a secure school workspace.")

    cards = [
        ("For learners", ["Assignments and revision", "Scores and feedback", "STEM practical activities", "Skills and study resources"]),
        ("For teachers", ["Create and publish assignments", "Auto-mark objective questions", "Review submissions", "Class and learner performance reports"]),
        ("For school leaders", ["Admission-number onboarding", "School identity and branding", "Role-based access", "Platform activity visibility"]),
    ]
    card_y = height - 480
    card_gap = 12
    card_w = (width - 76 - card_gap * 2) / 3
    for index, (title, items) in enumerate(cards):
        x = 38 + index * (card_w + card_gap)
        pdf.setFillColor(white)
        pdf.setStrokeColor(HexColor(LINE))
        pdf.roundRect(x, card_y, card_w, 190, 18, fill=1, stroke=1)
        pdf.setFillColor(HexColor(LIME))
        pdf.roundRect(x + 14, card_y + 146, 30, 28, 9, fill=1, stroke=0)
        pdf.setFillColor(HexColor(NAVY))
        pdf.setFont("Helvetica-Bold", 12)
        pdf.drawString(x + 14, card_y + 125, title)
        y = card_y + 98
        for item in items:
            pdf.setFillColor(HexColor(LIME))
            pdf.circle(x + 19, y + 3, 2.6, fill=1, stroke=0)
            y = draw_pdf_text(pdf, item, x + 29, y + 7, card_w - 42, "Helvetica", 9.2, SLATE, 13) - 7

    pdf.setFillColor(white)
    pdf.setStrokeColor(HexColor(LINE))
    pdf.roundRect(38, 135, width - 76, 170, 20, fill=1, stroke=1)
    pdf.setFillColor(HexColor(NAVY))
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(58, 272, "How a school rollout works")
    steps = [
        "1. School provides learner and staff onboarding records.",
        "2. Learners access a branded workspace using school-issued credentials.",
        "3. Teachers publish assignments and review performance by class or stream.",
        "4. School leaders use reports to support academic planning and parent discussions.",
    ]
    y = 242
    for step in steps:
        y = draw_pdf_text(pdf, step, 58, y, width - 230, "Helvetica", 10.3, SLATE, 15) - 7
    draw_qr_pdf(pdf, "https://assignmenthub.co.ke", width - 144, 171, 92)
    pdf.setFillColor(HexColor(SLATE))
    pdf.setFont("Helvetica", 8)
    pdf.drawCentredString(width - 98, 159, "View Assignment Hub")

    pdf.setFillColor(HexColor(NAVY))
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(38, 94, "Request a school demonstration")
    pdf.setFont("Helvetica", 10.5)
    pdf.drawString(38, 74, f"WhatsApp {PHONE}   |   {WEBSITE}")
    pdf.setFillColor(HexColor(MUTED))
    pdf.setFont("Helvetica", 7.5)
    pdf.drawRightString(width - 38, 42, "Product information for school review. No partnership or endorsement is implied.")
    pdf.save()
    return output


def paste_logo(base: Image.Image, box: tuple[int, int, int, int]) -> None:
    logo = cropped_logo().convert("RGBA")
    x1, y1, x2, y2 = box
    target_w = x2 - x1
    target_h = y2 - y1
    logo.thumbnail((target_w, target_h), Image.Resampling.LANCZOS)
    x = x1 + (target_w - logo.width) // 2
    y = y1 + (target_h - logo.height) // 2
    base.alpha_composite(logo, (x, y))


def draw_wrapped_pil(draw: ImageDraw.ImageDraw, text: str, xy: tuple[int, int], max_width: int, font: ImageFont.FreeTypeFont, fill: str, line_gap: int = 8) -> int:
    x, y = xy
    words = text.split()
    lines: list[str] = []
    line: list[str] = []
    for word in words:
        candidate = " ".join([*line, word])
        if line and draw.textbbox((0, 0), candidate, font=font)[2] > max_width:
            lines.append(" ".join(line))
            line = [word]
        else:
            line.append(word)
    if line:
        lines.append(" ".join(line))
    for text_line in lines:
        draw.text((x, y), text_line, font=font, fill=fill)
        y += font.size + line_gap
    return y


def build_status_ad() -> Path:
    output = OUT / "Assignment-Hub-WhatsApp-Status-Ad.png"
    width, height = 1080, 1920
    base = Image.new("RGBA", (width, height), NAVY)
    hero = ImageOps.fit(Image.open(HERO).convert("RGB"), (width, 850), method=Image.Resampling.LANCZOS, centering=(0.66, 0.5)).convert("RGBA")
    base.alpha_composite(hero, (0, 0))
    overlay = Image.new("RGBA", (width, 850), (16, 24, 32, 35))
    base.alpha_composite(overlay, (0, 0))
    draw = ImageDraw.Draw(base)

    draw.rounded_rectangle((58, 50, 650, 180), radius=28, fill=WHITE)
    paste_logo(base, (82, 72, 626, 160))
    draw.rounded_rectangle((60, 786, 190, 798), radius=6, fill=LIME)

    headline = "Schoolwork that fits in your pocket."
    y = draw_wrapped_pil(draw, headline, (60, 900), 900, pil_font(76, bold=True), WHITE, 10)
    y += 24
    y = draw_wrapped_pil(draw, "Assignments, feedback, STEM Labs, practical skills and study resources in one workspace.", (60, y), 920, pil_font(34), "#D7E0E7", 14)

    features = ["Complete assignments on your phone", "Get scores and feedback", "Learn Python, robotics and more"]
    y += 38
    for feature in features:
        draw.ellipse((62, y + 11, 80, y + 29), fill=LIME)
        draw.text((100, y), feature, font=pil_font(31, semibold=True), fill=WHITE)
        y += 66

    draw.rounded_rectangle((60, 1500, 1020, 1668), radius=34, fill=LIME)
    draw.text((100, 1534), "Independent access", font=pil_font(28, semibold=True), fill=NAVY)
    draw.text((100, 1581), "KSh 300 / month", font=pil_font(48, bold=True), fill=NAVY)
    draw.text((60, 1719), WEBSITE, font=pil_font(34, bold=True), fill=WHITE)
    draw.text((60, 1772), f"WhatsApp {PHONE}", font=pil_font(30), fill="#D7E0E7")

    qr_path = TMP / "register-qr.png"
    qr_png(REGISTER_URL, qr_path, 500)
    qr_image = Image.open(qr_path).convert("RGBA").resize((150, 150), Image.Resampling.LANCZOS)
    draw.rounded_rectangle((845, 1706, 1025, 1886), radius=24, fill=WHITE)
    base.alpha_composite(qr_image, (860, 1721))

    base.convert("RGB").save(output, quality=96)
    return output


def build_square_ad() -> Path:
    output = OUT / "Assignment-Hub-Social-Ad-Square.png"
    width = height = 1080
    base = Image.new("RGBA", (width, height), BACKGROUND)
    draw = ImageDraw.Draw(base)
    draw.rectangle((0, 0, 1080, 400), fill=NAVY)
    draw.rounded_rectangle((58, 48, 640, 170), radius=28, fill=WHITE)
    paste_logo(base, (82, 70, 615, 150))
    draw.rounded_rectangle((60, 342, 180, 354), radius=6, fill=LIME)
    y = draw_wrapped_pil(draw, "Assignments. Feedback. Skills. One workspace.", (60, 445), 930, pil_font(64, bold=True), NAVY, 12)
    y += 25
    y = draw_wrapped_pil(draw, "Built for Kenyan learners and schools. Works on phones, tablets and computers.", (60, y), 910, pil_font(32), SLATE, 12)

    chips = ["STEM Labs", "Study resources", "Instant feedback", "Practical skills"]
    chip_x, chip_y = 60, 730
    for chip in chips:
        chip_font = pil_font(25, semibold=True)
        chip_width = draw.textbbox((0, 0), chip, font=chip_font)[2] + 54
        if chip_x + chip_width > 1000:
            chip_x = 60
            chip_y += 75
        draw.rounded_rectangle((chip_x, chip_y, chip_x + chip_width, chip_y + 56), radius=22, fill="#E7F7B2")
        draw.text((chip_x + 27, chip_y + 11), chip, font=chip_font, fill=NAVY)
        chip_x += chip_width + 16

    draw.rounded_rectangle((60, 924, 1020, 1027), radius=28, fill=NAVY)
    draw.text((95, 947), WEBSITE, font=pil_font(30, bold=True), fill=WHITE)
    draw.text((620, 949), f"WhatsApp {PHONE}", font=pil_font(27), fill=LIME)
    base.convert("RGB").save(output, quality=96)
    return output


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def build_email_template() -> Path:
    output = OUT / "Assignment-Hub-School-Outreach-Email-Template.docx"
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("334155")
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.12

    logo_paragraph = document.add_paragraph()
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    logo_paragraph.add_run().add_picture(str(LOGO), width=Inches(2.35))

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run("School outreach email template")
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor.from_string("101820")

    subtitle = document.add_paragraph("A concise introduction for principals, directors and academic leaders")
    subtitle.runs[0].font.color.rgb = RGBColor.from_string("64748B")

    table = document.add_table(rows=2, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.15)
    table.columns[1].width = Inches(5.7)
    labels = [("To", "[Principal or Director Name]"), ("Subject", "A simpler way to manage assignments and learner feedback at [School Name]")]
    for row, (label, value) in zip(table.rows, labels):
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(row.cells[0], "101820")
        set_cell_shading(row.cells[1], "F8FAFC")
        row.cells[0].paragraphs[0].add_run(label).font.color.rgb = RGBColor(255, 255, 255)
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        row.cells[1].paragraphs[0].add_run(value)

    document.add_paragraph()
    body = [
        "Dear [Title and Name],",
        "I am writing to introduce Assignment Hub, a Kenyan web-based learning platform designed to reduce the time and cost schools spend distributing, collecting and reviewing hard-copy assignments.",
        "Learners use a secure school workspace on a phone, tablet or computer. They can complete assignments, receive released scores and feedback, access STEM practical activities, build technical skills and use approved study resources. Teachers can publish work, review submissions and follow class or stream performance without changing the school's existing academic structure.",
        "For school leaders, Assignment Hub supports learner onboarding using school codes and admission numbers, role-based access, school identity, subscription management and platform activity visibility.",
        "I would value 20 minutes to demonstrate the learner, teacher and school administration experience using a pilot class or department of your choice.",
        "Would [suggested date and time] be convenient for a short online or in-person demonstration?",
        "Kind regards,",
        "Frank",
        f"Assignment Hub | Andolih EdTech Studios\nWhatsApp: {PHONE}\nWebsite: {WEBSITE}",
    ]
    for index, paragraph_text in enumerate(body):
        paragraph = document.add_paragraph()
        for line_index, line in enumerate(paragraph_text.split("\n")):
            if line_index:
                paragraph.add_run().add_break()
            run = paragraph.add_run(line)
            if index in {0, 7}:
                run.font.bold = True

    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(14)
    run = heading.add_run("Follow-up email after 4 to 5 working days")
    run.font.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string("101820")

    follow_up = (
        "Dear [Title and Name],\n\n"
        "I am following up on my note about Assignment Hub. The platform gives schools one workspace for assignments, feedback, practical learning, approved resources and performance visibility across phones and computers.\n\n"
        "I would be glad to demonstrate the system using a short school scenario and answer questions from your academic or ICT team. Please let me know a suitable time, or reach me on WhatsApp at "
        f"{PHONE}.\n\nKind regards,\nFrank"
    )
    for block in follow_up.split("\n\n"):
        document.add_paragraph(block)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(f"Assignment Hub  |  {WEBSITE}  |  {PHONE}")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string("64748B")

    document.core_properties.title = "Assignment Hub School Outreach Email Template"
    document.core_properties.author = "Andolih EdTech Studios"
    document.save(output)
    return output


def build_plain_email() -> Path:
    output = OUT / "Assignment-Hub-School-Outreach-Email-Template.txt"
    output.write_text(
        "Subject: A simpler way to manage assignments and learner feedback at [School Name]\n\n"
        "Dear [Title and Name],\n\n"
        "I am writing to introduce Assignment Hub, a Kenyan web-based learning platform designed to reduce the time and cost schools spend distributing, collecting and reviewing hard-copy assignments.\n\n"
        "Learners use a secure school workspace on a phone, tablet or computer. They can complete assignments, receive released scores and feedback, access STEM practical activities, build technical skills and use approved study resources. Teachers can publish work, review submissions and follow class or stream performance without changing the school's existing academic structure.\n\n"
        "For school leaders, Assignment Hub supports learner onboarding using school codes and admission numbers, role-based access, school identity, subscription management and platform activity visibility.\n\n"
        "I would value 20 minutes to demonstrate the learner, teacher and school administration experience using a pilot class or department of your choice. Would [suggested date and time] be convenient for a short online or in-person demonstration?\n\n"
        "Kind regards,\nFrank\nAssignment Hub | Andolih EdTech Studios\n"
        f"WhatsApp: {PHONE}\nWebsite: {WEBSITE}\n",
        encoding="utf-8",
    )
    return output


def main() -> None:
    outputs = [
        build_student_pamphlet(),
        build_nyangorri_snapshot(),
        build_status_ad(),
        build_square_ad(),
        build_plain_email(),
    ]
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
